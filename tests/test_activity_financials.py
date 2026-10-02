"""Activity financials and all-tier source metadata contract coverage."""

import copy
import json
import sys
from unittest.mock import MagicMock

import httpx
import pytest

from handelsregister import (
    ActivityStatement,
    Company,
    FinancialAccount,
    FinancialProvenance,
    FinancialStatement,
    Handelsregister,
    OrganizationFeature,
)
from handelsregister import cli


@pytest.fixture
def financial_response():
    provenance = {
        "statement_type": "Jahresabschluss",
        "period_start": "2023-01-01",
        "period_end": "2023-12-31",
        "exempt_subsidiary": False,
        "parent_organization": None,
    }
    activity_provenance = {**provenance, "statement_type": "Tätigkeitsabschluss"}
    activity = {
        "name": {
            "en": "Electricity distribution",
            "de": "Elektrizitätsverteilung",
            "in_report": "Strom Netz",
        }
    }
    balance = [
        {
            "name": {"en": "Assets", "de": "Aktivseite"},
            "value": 7668374.41,
            "children": [
                {
                    "name": {"en": "Fixed Assets", "in_report": "A. Anlagevermögen"},
                    "value": 5842910.1,
                    "children": [
                        {
                            "name": {"en": "Land"},
                            "value": 0,
                            "children": [],
                            "future_account_field": "preserved",
                        }
                    ],
                }
            ],
        }
    ]
    pnl = [
        {
            "name": {"en": "Revenue", "de": "Umsatzerlöse"},
            "value": 100,
            "children": [
                {"name": {"en": "Distribution revenue"}, "value": 90, "children": []}
            ],
        },
        {"name": {"en": "Net income"}, "value": -5, "children": []},
    ]
    return {
        "name": "Example energy company",
        "balance_sheet_accounts": [
            {
                "year": 2024,
                "balance_sheet_accounts": [],
                "_provenance": {
                    **provenance,
                    "period_start": "2024-01-01",
                    "period_end": "2024-12-31",
                },
            },
            {
                "year": 2023,
                "balance_sheet_accounts": balance,
                "_provenance": provenance,
                "activity_statements": [
                    {
                        "activity": activity,
                        "balance_sheet_accounts": balance,
                        "_provenance": activity_provenance,
                    }
                ],
            },
        ],
        "profit_and_loss_account": [
            {
                "year": 2024,
                "profit_and_loss_accounts": [],
                "_provenance": {
                    **provenance,
                    "period_start": "2024-01-01",
                    "period_end": "2024-12-31",
                },
            },
            {
                "year": 2023,
                "profit_and_loss_accounts": pnl,
                "_provenance": provenance,
                "activity_statements": [
                    {
                        "activity": {
                            "name": {**activity["name"], "in_report": "Stromnetz"}
                        },
                        "profit_and_loss_accounts": pnl,
                        "_provenance": activity_provenance,
                    }
                ],
            },
        ],
    }


def company_from(data):
    client = MagicMock(spec=Handelsregister)
    client.fetch_organization.return_value = data
    return Company("example", client=client)


def test_legacy_api_response_remains_compatible(
    sample_organization_response, monkeypatch, capsys
):
    """Older servers have neither activities, provenance nor deep shareholders."""
    data = sample_organization_response
    assert "shareholders_deep" not in data
    for feature in (
        "financial_kpi",
        "balance_sheet_accounts",
        "profit_and_loss_account",
    ):
        assert all(
            "_provenance" not in row and "activity_statements" not in row
            for row in data.get(feature, [])
        )
    requests = []

    def handler(request):
        requests.append(request)
        return httpx.Response(200, json=data)

    real_client = httpx.Client
    monkeypatch.setattr(
        httpx,
        "Client",
        lambda **kwargs: real_client(transport=httpx.MockTransport(handler), **kwargs),
    )
    features = [
        "financial_kpi",
        "balance_sheet_accounts",
        "profit_and_loss_account",
        "shareholders",
    ]
    client = Handelsregister(api_key="test")
    company = Company("legacy-company", client=client, features=features)
    assert requests[0].url.params.get_list("feature") == features
    assert company.data == data
    assert company.financial_kpi == data["financial_kpi"]
    assert company.balance_sheet_accounts == data["balance_sheet_accounts"]
    assert company.profit_and_loss_account == data["profit_and_loss_account"]
    assert company.capital == data.get("capital", "")
    assert company.shareholders.as_dict() == data.get("shareholders", {})
    assert not company.shareholders_deep
    for statement in company.balance_sheet_entries + company.profit_and_loss_entries:
        assert not statement.provenance
        assert statement.activity_statements == []
        assert company.get_activity_balance_sheets_for_year(statement.year) == []
        assert company.get_activity_profit_and_loss_for_year(statement.year) == []
    for rich in (True, False):
        monkeypatch.setattr(cli, "RICH_AVAILABLE", rich)
        cli._display_result(client, data)
        output = capsys.readouterr().out
        assert company.name in output
        assert "Activity Balance sheet" not in output
        assert "Deep Shareholders" not in output
    flattened = client._flatten_result(data)
    assert not any(
        key.endswith("_provenance") or key.endswith("_activity_statements")
        for key in flattened
    )


def test_activity_helpers_and_typed_account_trees(financial_response):
    before = copy.deepcopy(financial_response)
    company = company_from(financial_response)
    balance = company.get_balance_sheet_entry_for_year(2023)
    pnl = company.get_profit_and_loss_entry_for_year(2023)
    assert isinstance(balance, FinancialStatement)
    assert balance.provenance.statement_type == "Jahresabschluss"
    assert balance.provenance.exempt_subsidiary is False
    assert balance.balance_sheet_entries[0].value == 7668374.41
    activity = company.get_activity_balance_sheets_for_year(2023)[0]
    assert isinstance(activity, ActivityStatement)
    assert activity.name_text() == "Electricity distribution"
    assert activity.name_text("de") == "Elektrizitätsverteilung"
    assert activity.name_in_report == "Strom Netz"
    assert activity.provenance.statement_type == "Tätigkeitsabschluss"
    assert activity.provenance.period_end == "2023-12-31"
    assert activity.profit_and_loss_entries == []
    root = activity.balance_sheet_entries[0]
    assert isinstance(root, FinancialAccount)
    assert root.name_text("de") == "Aktivseite"
    assert root.children[0].value == 5842910.1
    nodes = list(root.walk())
    assert [node.name_text() for node in nodes] == ["Assets", "Fixed Assets", "Land"]
    assert nodes[-1].value == 0
    assert nodes[-1].as_dict()["future_account_field"] == "preserved"
    activity_pnl = company.get_activity_profit_and_loss_for_year(2023)[0]
    assert (
        activity_pnl.name_in_report == "Stromnetz"
    )  # Labels need not match between reports.
    assert activity_pnl.balance_sheet_entries == []
    assert activity_pnl.profit_and_loss_entries[-1].value == -5
    assert activity_pnl.provenance.as_dict() == activity.provenance.as_dict()
    assert pnl.profit_and_loss_entries[0].children[0].value == 90
    # Existing raw interfaces, including their metadata, are preserved.
    assert company.get_balance_sheet_for_year(2023) is balance.raw
    assert company.get_profit_and_loss_for_year(2023) is pnl.raw
    assert financial_response == before


@pytest.mark.parametrize("year", [2024, 1900])
def test_years_without_activities_and_missing_years(financial_response, year):
    company = company_from(financial_response)
    assert company.get_activity_balance_sheets_for_year(year) == []
    assert company.get_activity_profit_and_loss_for_year(year) == []
    if year == 1900:
        assert company.get_balance_sheet_entry_for_year(year) is None
        assert company.get_profit_and_loss_entry_for_year(year) is None


@pytest.mark.parametrize("activities", [None, [], "invalid"])
def test_optional_null_and_malformed_collections(activities):
    statement = FinancialStatement.from_payload(
        {"year": 2023, "activity_statements": activities}
    )
    assert statement.activity_statements == []
    assert statement.provenance.as_dict() == {}
    assert not statement.provenance
    activity = ActivityStatement.from_payload(
        {
            "activity": {"name": {"in_report": "Original activity"}},
            "balance_sheet_accounts": None,
        }
    )
    assert activity.name_text() == "Original activity"
    assert activity.balance_sheet_entries == []
    assert activity.profit_and_loss_entries == []
    assert (
        FinancialAccount.from_payload({"name": "Legacy", "children": None}).children
        == []
    )
    assert FinancialAccount.from_accounts({"revenue": 0})[0].value == 0


@pytest.mark.parametrize("plan", ["free", "basic", "pro", "max"])
def test_provenance_on_every_tier_without_a_new_feature_or_plan_preflight(
    financial_response, monkeypatch, plan
):
    data = copy.deepcopy(financial_response)
    if plan != "max":
        for feature in ("balance_sheet_accounts", "profit_and_loss_account"):
            for row in data[feature]:
                row.pop("activity_statements", None)
    parent = {
        "entity_id": "parent-id",
        "name": "Parent company",
        "future_parent_field": True,
    }
    source = data["balance_sheet_accounts"][1]["_provenance"]
    source.update(
        {
            "exempt_subsidiary": True,
            "parent_organization": parent,
            "future_source_field": "kept",
        }
    )
    requests = []

    def handler(request):
        requests.append(request)
        return httpx.Response(200, json=data)

    real_client = httpx.Client
    monkeypatch.setattr(
        httpx,
        "Client",
        lambda **kwargs: real_client(transport=httpx.MockTransport(handler), **kwargs),
    )
    company = Company(
        "energy-company",
        client=Handelsregister(api_key="test"),
        features=[
            OrganizationFeature.BALANCE_SHEET_ACCOUNTS,
            OrganizationFeature.PROFIT_AND_LOSS_ACCOUNT,
        ],
    )
    assert len(requests) == 1
    assert requests[0].url.params.get_list("feature") == [
        "balance_sheet_accounts",
        "profit_and_loss_account",
    ]
    provenance = company.get_balance_sheet_entry_for_year(2023).provenance
    assert isinstance(provenance, FinancialProvenance)
    assert provenance.exempt_subsidiary is True
    assert provenance.parent_organization == parent
    assert provenance.as_dict()["future_source_field"] == "kept"
    assert bool(company.get_activity_balance_sheets_for_year(2023)) == (plan == "max")


@pytest.mark.parametrize("rich", [True, False])
def test_cli_selected_year_shows_separate_activity_accounts_and_sources(
    financial_response, monkeypatch, capsys, rich
):
    monkeypatch.setattr(cli, "RICH_AVAILABLE", rich)
    monkeypatch.setenv("COLUMNS", "220")
    cli._display_result(
        Handelsregister(api_key="test"), financial_response, financial_year=2023
    )
    out = capsys.readouterr().out
    assert "Activity Balance sheet 2023: Electricity distribution" in out
    assert "Activity P&L 2023: Electricity distribution" in out
    assert "7,668,374.41 €" in out
    assert "Land" in out
    assert "0.00 €" in out
    assert "-5.00 €" in out
    assert "Tätigkeitsabschluss" in out
    assert "2023-01-01" in out
    assert "2024-01-01" not in out
    assert "_provenance" not in out


def test_cli_year_selection_is_local_and_json_stays_complete(
    financial_response, monkeypatch, capsys
):
    requests = []

    def fetch(self, **kwargs):
        requests.append(kwargs)
        return financial_response

    monkeypatch.setattr(Handelsregister, "fetch_organization", fetch)
    monkeypatch.setenv("HANDELSREGISTER_API_KEY", "test")
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "prog",
            "fetch",
            "json",
            "energy-company",
            "--feature",
            "balance_sheet_accounts",
            "--financial-year",
            "2023",
        ],
    )
    cli.main()
    assert json.loads(capsys.readouterr().out) == financial_response
    assert requests[0] == {
        "q": "energy-company",
        "features": ["balance_sheet_accounts"],
        "ai_search": "on-default",
        "realtime_mode": None,
    }
    year, _, tables = cli._financial_display(financial_response, None)
    assert year == 2024
    assert tables == []


@pytest.mark.parametrize("output_type", ["json", "csv", "xlsx"])
def test_export_preserves_both_activity_trees_and_provenance(
    financial_response, tmp_path, monkeypatch, output_type
):
    import pandas as pd

    client = Handelsregister(api_key="test")
    monkeypatch.setattr(
        client, "fetch_organization", lambda **kwargs: financial_response
    )
    source = tmp_path / "companies.json"
    source.write_text(json.dumps([{"name": "Energy company"}]))
    target = tmp_path / ("output." + output_type)
    client.enrich(
        file_path=str(source),
        input_type="json",
        output_file=str(target),
        output_type=output_type,
        query_properties={"name": "name"},
        params={"features": ["balance_sheet_accounts", "profit_and_loss_account"]},
    )
    if output_type == "json":
        assert (
            json.loads(target.read_text())[0]["_handelsregister_result"]
            == financial_response
        )
    else:
        frame = pd.read_csv(target) if output_type == "csv" else pd.read_excel(target)
        for feature in ("balance_sheet_accounts", "profit_and_loss_account"):
            metadata = json.loads(frame.iloc[0]["hr_" + feature + "_provenance"])
            assert [row["_provenance"] for row in metadata] == [
                row["_provenance"] for row in financial_response[feature]
            ]
            activities = json.loads(
                frame.iloc[0]["hr_" + feature + "_activity_statements"]
            )
            assert activities == [
                {
                    "year": 2023,
                    "activity_statements": financial_response[feature][1][
                        "activity_statements"
                    ],
                }
            ]


def test_flat_account_summaries_handle_string_labels_and_null_children():
    client = Handelsregister(api_key="test")
    assert client._flatten_account(
        {"name": "Revenue", "value": 0, "children": None}
    ) == ["Revenue: 0"]


def test_excel_export_reconstructs_oversized_activity_json(
    financial_response, tmp_path, monkeypatch
):
    import pandas as pd

    row = financial_response["balance_sheet_accounts"][1]
    row["activity_statements"][0]["balance_sheet_accounts"][0]["name"]["in_report"] = (
        "Long activity report " * 5000
    )
    client = Handelsregister(api_key="test")
    monkeypatch.setattr(
        client, "fetch_organization", lambda **kwargs: financial_response
    )
    source = tmp_path / "input.json"
    source.write_text(json.dumps([{"name": "Energy company"}]))
    target = tmp_path / "output.xlsx"
    client.enrich(
        file_path=str(source),
        input_type="json",
        output_file=str(target),
        output_type="xlsx",
        query_properties={"name": "name"},
    )
    sheets = pd.read_excel(target, sheet_name=None)
    column = "hr_balance_sheet_accounts_activity_statements"
    reference = json.loads(sheets["Sheet1"].iloc[0][column])[
        "_handelsregister_excel_overflow"
    ]
    assert reference["sheet"] == "Long values"
    assert reference["row"] == 2
    assert reference["column"] == column
    chunks = sheets["Long values"]
    chunks = chunks[
        (chunks["row"] == reference["row"]) & (chunks["column"] == column)
    ].sort_values("part")
    assert len(chunks) == reference["chunks"]
    assert all(len(value) <= 16000 for value in chunks["value"])
    assert json.loads("".join(chunks["value"])) == [
        {"year": 2023, "activity_statements": row["activity_statements"]}
    ]


@pytest.mark.parametrize(
    "value",
    ["x" * 32767, "x" * 32768, "😀" * 20000, "x" * 16000 + "=1+1" + "x" * 18000],
    ids=["cell-limit", "cell-overflow", "unicode-overflow", "literal-chunks"],
)
def test_excel_limits_unicode_and_literal_chunks(tmp_path, value):
    import pandas as pd
    from openpyxl import load_workbook

    frame = pd.DataFrame({"note": [value, "Small value"]}, index=[10, 20])
    target = tmp_path / "output.xlsx"
    Handelsregister._write_enrichment_excel(frame, str(target))
    workbook = load_workbook(target, data_only=False)
    assert workbook["Sheet1"]["A3"].value == "Small value"
    if len(value.encode("utf-16-le")) // 2 <= 32767:
        assert workbook["Sheet1"]["A2"].value == value
        assert "Long values" not in workbook.sheetnames
    else:
        reference = json.loads(workbook["Sheet1"]["A2"].value)[
            "_handelsregister_excel_overflow"
        ]
        cells = [row[3] for row in workbook[reference["sheet"]].iter_rows(min_row=2)]
        assert "".join(cell.value for cell in cells) == value
        assert all(cell.data_type == "s" for cell in cells)
        assert all(len(cell.value.encode("utf-16-le")) // 2 <= 32767 for cell in cells)
    assert frame.iloc[0]["note"] == value
